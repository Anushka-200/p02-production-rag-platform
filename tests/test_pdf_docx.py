import io
import re
from pathlib import Path

import docx
import pymupdf
import pytest

from app.ingest.common import join_wrapped_lines, table_to_markdown
from app.ingest.models import ParseError
from app.ingest.parsers import parse, parse_docx, parse_pdf
from scripts.make_sample_docx import build

CSS = (
    "h1{font-size:20pt} h2{font-size:15pt} p,li,td,th{font-size:10pt} "
    "table{border-collapse:collapse} td,th{border:1px solid black;padding:3px}"
)
PAGE1 = (
    "<h1>Acme Benefits Guide</h1><p>Welcome to the guide. It explains your benefits.</p>"
    "<h2>Time off</h2><p>Employees get twenty vacation days per year. Unused days do not carry "
    "over except in special cases described below.</p>"
)
PAGE2 = (
    "<h2>Expense limits</h2><p>Meals are capped as shown.</p>"
    "<table><tr><th>Item</th><th>Cap</th></tr><tr><td>Meal</td><td>50</td></tr>"
    "<tr><td>Hotel</td><td>200</td></tr></table><p>Receipts are required.</p>"
)


def make_pdf(*pages_html: str) -> bytes:
    doc = pymupdf.open()
    for html in pages_html:
        page = doc.new_page()
        page.insert_htmlbox(page.rect + (36, 36, -36, -36), html, css=CSS)
    return doc.tobytes()  # type: ignore[no-any-return]


# ---- shared helpers ---------------------------------------------------------------------------
def test_table_to_markdown_escapes_pipes_flattens_newlines_and_pads_rows() -> None:
    md = table_to_markdown([["Item", "Note"], ["a|b", "line1\nline2"], ["short"], ["", ""]])
    assert md.splitlines() == [
        "| Item | Note |",
        "|---|---|",
        "| a\\|b | line1 line2 |",
        "| short |  |",  # padded; the fully empty row is dropped
    ]


def test_table_to_markdown_empty_input() -> None:
    assert table_to_markdown([]) == "" and table_to_markdown([["", ""]]) == ""


def test_join_wrapped_lines_repairs_hyphenation_only_before_lowercase() -> None:
    assert join_wrapped_lines(["an em-", "ployee of the"]) == "an employee of the"
    assert join_wrapped_lines(["see Pub. 15-", "B for more"]) == "see Pub. 15- B for more"
    assert join_wrapped_lines(["first line", "second line"]) == "first line second line"


# ---- PDF: generated file with known content -------------------------------------------------
def test_pdf_pages_headings_and_title() -> None:
    doc = parse_pdf(make_pdf(PAGE1, PAGE2), "guide.pdf")
    assert doc.title == "Acme Benefits Guide" and doc.metadata["pages"] == 2
    got = [(s.page, s.heading_path) for s in doc.sections]
    assert got == [
        (1, ("Acme Benefits Guide",)),
        (1, ("Acme Benefits Guide", "Time off")),
        (2, ("Acme Benefits Guide", "Expense limits")),
    ]
    assert doc.sections[2].anchor == "p2-expense-limits"


def test_pdf_ligatures_are_normalised_so_keyword_search_works() -> None:
    doc = parse_pdf(make_pdf(PAGE1), "guide.pdf")
    text = " ".join(s.text for s in doc.sections) + doc.title
    assert "benefits" in text and "ﬁ" not in text  # no "ﬁ" ligature left


def test_pdf_table_becomes_one_pipe_table_not_duplicated_text() -> None:
    doc = parse_pdf(make_pdf(PAGE1, PAGE2), "guide.pdf")
    text = doc.sections[2].text
    assert "| Item | Cap |" in text and "| Hotel | 200 |" in text
    assert text.count("Hotel") == 1  # the table's text is not also kept as plain text
    assert doc.metadata["tables"] == 1


def test_pdf_with_no_text_layer_raises_a_clear_error() -> None:
    blank = pymupdf.open()
    blank.new_page()
    with pytest.raises(ParseError, match="no extractable text"):
        parse_pdf(blank.tobytes(), "scan.pdf")


def test_corrupt_pdf_raises_parse_error() -> None:
    with pytest.raises(ParseError, match="cannot open PDF"):
        parse_pdf(b"this is not a pdf", "bad.pdf")


# ---- DOCX -------------------------------------------------------------------------------------
def docx_bytes() -> bytes:
    buf = io.BytesIO()
    build().save(buf)
    return buf.getvalue()


def test_docx_headings_lists_and_table() -> None:
    d = parse_docx(docx_bytes(), "policy.docx")
    assert d.title == "Acme Corp Benefits Policy"
    paths = [s.heading_path for s in d.sections]
    assert paths == [
        ("Acme Corp Benefits Policy",),
        ("Acme Corp Benefits Policy", "Time off"),
        ("Acme Corp Benefits Policy", "Time off", "Parental leave"),
        ("Acme Corp Benefits Policy", "Expense limits"),
    ]
    time_off = d.sections[1].text
    assert "- Sick leave is paid for up to 10 days." in time_off  # list items keep their bullet
    expenses = d.sections[3].text
    assert "| Item | US | UK |" in expenses and "| Hotel | 200 | 150 |" in expenses
    assert all(s.page is None for s in d.sections)  # Word files have no fixed pages


def test_docx_merged_cells_are_not_repeated() -> None:
    d = docx.Document()
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).merge(t.cell(0, 2)).text = "Benefits overview"
    for c, v in enumerate(["a", "b", "c"]):
        t.cell(1, c).text = v
    buf = io.BytesIO()
    d.save(buf)
    text = parse_docx(buf.getvalue(), "m.docx").sections[0].text
    assert text.count("Benefits overview") == 1


def test_corrupt_docx_raises_parse_error() -> None:
    with pytest.raises(ParseError, match="cannot open DOCX"):
        parse_docx(b"not a docx", "bad.docx")


def test_sample_docx_file_in_repo_parses() -> None:
    path = Path("data/samples/acme-benefits-policy.docx")
    assert parse("docx", path.read_bytes(), path.name).sections


# ---- real IRS / OPM PDFs (skipped when the downloads are missing) ------------------------------
HR = Path(__file__).resolve().parents[1] / "data" / "raw" / "hr-pdfs"


@pytest.mark.skipif(
    not (HR / "irs-p15b-fringe-benefits.pdf").exists(), reason="PDFs not downloaded"
)
def test_real_irs_pdf_has_pages_headings_table_and_repaired_hyphens() -> None:
    doc = parse("pdf", (HR / "irs-p15b-fringe-benefits.pdf").read_bytes(), "p15b.pdf")
    assert doc.metadata["pages"] == 36 and max(s.page or 0 for s in doc.sections) == 36
    assert any("Fringe Benefit Exclusion Rules" in " > ".join(s.heading_path) for s in doc.sections)
    table_pages = {s.page for s in doc.sections if "|---|" in s.text}
    assert 6 in table_pages  # Table 2-1 is on page 6 of the real file
    text = "\n".join(s.text for s in doc.sections)
    assert not re.search(r"\bployee", text)  # no orphaned second half of "em- ployee"


@pytest.mark.skipif(not HR.exists(), reason="PDFs not downloaded")
def test_every_real_pdf_parses() -> None:
    for pdf in HR.glob("*.pdf"):
        doc = parse("pdf", pdf.read_bytes(), pdf.name)
        assert doc.sections and all(s.page for s in doc.sections), pdf.name
