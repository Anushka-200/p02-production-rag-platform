"""Writes the synthetic sample DOCX (our own text, no licence issue) used to test the DOCX parser.

Run: uv run python -m scripts.make_sample_docx
"""

from pathlib import Path

import docx


def build() -> docx.document.Document:
    d = docx.Document()
    d.core_properties.title = "Acme Corp Benefits Policy"
    d.add_heading("Acme Corp Benefits Policy", level=1)
    d.add_paragraph("This policy applies to all full-time employees of Acme Corp.")
    d.add_heading("Time off", level=2)
    d.add_paragraph("Employees receive 20 vacation days per year.")
    d.add_paragraph("Sick leave is paid for up to 10 days.", style="List Bullet")
    d.add_paragraph("A doctor's note is needed after 3 days.", style="List Bullet")
    d.add_heading("Parental leave", level=3)
    d.add_paragraph("Primary caregivers receive 16 weeks of paid leave.")
    d.add_heading("Expense limits", level=2)
    table = d.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for row, values in enumerate(
        [("Item", "US", "UK"), ("Meal", "50", "40"), ("Hotel", "200", "150")]
    ):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    d.add_paragraph("Receipts are required for all claims.")
    return d


if __name__ == "__main__":
    out = Path("data/samples/acme-benefits-policy.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    build().save(out)
    print("wrote", out)
